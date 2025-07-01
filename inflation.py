import requests
import pandas as pd
import numpy as np
from datetime import datetime
import openpyxl
import matplotlib.pyplot as plt
from openpyxl.drawing.image import Image
from pathlib import Path
import shutil

# --- PARTE 1: EXTRAÇÃO DOS DADOS DO BCB ---
def get_bcb_data(series_code):
    """Obtém dados do BCB via API do SGS e retorna um DataFrame"""
    url = f"https://api.bcb.gov.br/dados/serie/bcdata.sgs.{series_code}/dados?formato=json"
    try:
        response = requests.get(url)
        response.raise_for_status()
        dados = response.json()
        df = pd.DataFrame(dados)
        df['data'] = pd.to_datetime(df['data'], dayfirst=True)
        df['valor'] = pd.to_numeric(df['valor'].str.replace(',', '.'))
        return df
    except Exception as e:
        print(f"Erro ao obter série {series_code}: {e}")
        return None

# Códigos das séries no BCB (SGS)
codigos = {
    'IPCA': 433,  # IPCA - Variação mensal (%)
    'IGPM': 189,  # IGP-M - Variação mensal (%)
}

dfs = {}
for nome, codigo in codigos.items():
    df = get_bcb_data(codigo)
    if df is not None:
        dfs[nome] = df[['data', 'valor']].rename(columns={'valor': nome})

# Unir os dados em um único DataFrame mensal
if len(dfs) == 2:
    df_final = pd.merge(dfs['IPCA'], dfs['IGPM'], on='data', how='outer')
    df_final = df_final.sort_values('data').reset_index(drop=True)
    # Formatar data para DD/MM/YYYY
    df_final['data'] = df_final['data'].dt.strftime('%d/%m/%Y')
    # Salvar em CSV mensal
    nome_arquivo = f"ipca_igpm_{datetime.now().strftime('%Y%m%d')}.csv"
    df_final.to_csv(
        nome_arquivo,
        index=False,
        sep=',',
        quotechar='\0',  # Caractere nulo para evitar aspas
        quoting=3,       # QUOTE_NONE (sem aspas)
        float_format='%.2f',
        encoding='utf-8'
    )
    print(f"✅ Arquivo mensal salvo: {nome_arquivo}")
    print(f"📊 Total de registros: {len(df_final)}")
else:
    print("❌ Falha ao obter dados de uma das séries.")
    exit(1)

# --- PARTE 2: CÁLCULO DOS TRIMESTRES ---
csv_in = nome_arquivo
df = pd.read_csv(csv_in, encoding='utf-8')
df['data'] = pd.to_datetime(df['data'], format='%d/%m/%Y', errors='coerce')
df = df.dropna(subset=['data'])

def get_trimestre(dt):
    trimestre = (dt.month - 1) // 3 + 1
    return f'{trimestre}Q{str(dt.year)[-2:]}'

df['trimestre'] = df['data'].apply(get_trimestre)

def acumulado_composto(serie):
    serie = pd.to_numeric(serie, errors='coerce').dropna()
    if len(serie) == 0:
        return np.nan
    fatores = (1 + serie/100)
    return (fatores.prod() - 1) * 100

trimestres_unicos = df['trimestre'].unique()
trimestres_unicos.sort()

resultados = []
for trimestre in trimestres_unicos:
    grupo = df[df['trimestre'] == trimestre]
    ipca = acumulado_composto(grupo['IPCA'])
    igpm = acumulado_composto(grupo['IGPM'])
    # Formatar com vírgula como separador decimal
    ipca_str = f"{ipca:.2f}".replace('.', ',') if not pd.isna(ipca) else ''
    igpm_str = f"{igpm:.2f}".replace('.', ',') if not pd.isna(igpm) else ''
    resultados.append({'trimestre': trimestre, 'IPCA': ipca_str, 'IGPM': igpm_str})

trimestral = pd.DataFrame(resultados)
trimestral['ano'] = trimestral['trimestre'].str.extract(r'(\d{2})$').astype(int)
trimestral['tri'] = trimestral['trimestre'].str.extract(r'^(\d)Q').astype(int)
trimestral = trimestral.sort_values(['ano', 'tri']).drop(['ano', 'tri'], axis=1).reset_index(drop=True)

# Substituir NaN por string vazia antes de salvar
trimestral = trimestral.replace({np.nan: ''})
# Salvar em novo CSV
csv_out = 'ipca_igpm_trimestres.csv'
trimestral.to_csv(csv_out, index=False, encoding='utf-8', sep=',')

print('\n--- DADOS MENSAIS ---')
print(df_final)
print('\n--- DADOS TRIMESTRAIS ---')
print(trimestral)

def add_charts_with_xlsxwriter(input_excel, output_excel):
    # Lê a primeira aba do consolidado
    df = pd.read_excel(input_excel, sheet_name=0)
    
    # Filtra apenas as linhas de SSS_Descontado e SSR_Descontado
    sss = df[df['Métrica'] == 'SSS_Descontado']
    ssr = df[df['Métrica'] == 'SSR_Descontado']
    
    # Cria um novo arquivo Excel com XlsxWriter
    with pd.ExcelWriter(output_excel, engine='xlsxwriter') as writer:
        df.to_excel(writer, sheet_name='Consolidado', index=False)
        workbook  = writer.book
        worksheet = writer.sheets['Consolidado']
        
        # Descobre o range dos dados
        n_trimestres = len(df.columns) - 2
        first_row = 1  # Cabeçalho está na linha 0
        sss_rows = sss.index.tolist()
        ssr_rows = ssr.index.tolist()
        
        # Adiciona gráfico de linhas para SSS_Descontado
        chart1 = workbook.add_chart({'type': 'line'})
        for row in sss_rows:
            chart1.add_series({
                'name':       ['Consolidado', row+1, 0],
                'categories': ['Consolidado', 0, 2, 0, n_trimestres+1],
                'values':     ['Consolidado', row+1, 2, row+1, n_trimestres+1],
            })
        chart1.set_title({'name': 'SSS_Descontado'})
        chart1.set_x_axis({'name': 'Trimestre'})
        chart1.set_y_axis({'name': '%'})
        worksheet.insert_chart('A40', chart1)
        
        # Adiciona gráfico de linhas para SSR_Descontado
        chart2 = workbook.add_chart({'type': 'line'})
        for row in ssr_rows:
            chart2.add_series({
                'name':       ['Consolidado', row+1, 0],
                'categories': ['Consolidado', 0, 2, 0, n_trimestres+1],
                'values':     ['Consolidado', row+1, 2, row+1, n_trimestres+1],
            })
        chart2.set_title({'name': 'SSR_Descontado'})
        chart2.set_x_axis({'name': 'Trimestre'})
        chart2.set_y_axis({'name': '%'})
        worksheet.insert_chart('J40', chart2)
        
    print(f"Arquivo com gráficos salvo em: {output_excel}")

# Exemplo de uso:
add_charts_with_xlsxwriter('data_treated/Consolidado.xlsx', 'data_treated/Consolidado_com_graficos.xlsx') 